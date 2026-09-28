import base64
import json
import mimetypes
import zipfile
import re
from pathlib import Path
import xml.etree.ElementTree as ET

import openpyxl


# ============================================================
# KONFIGURASI
# ============================================================

BASE_DIR = Path(__file__).resolve().parent

EXCEL_FILE = BASE_DIR / "SPAREPART BNPB.xlsx"
LOGO_FILE = BASE_DIR / "logo_bnpb.png"
OUTPUT_FILE = BASE_DIR / "index.html"


# ============================================================
# XML NAMESPACE
# ============================================================

NS = {
    "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "rel": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "pkgrel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}

RID_ATTR = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
EMBED_ATTR = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"


# ============================================================
# MEMBERSIHKAN DATA EXCEL
# ============================================================

def clean(value):
    if value is None:
        return ""

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    return str(value).strip()


# ============================================================
# IMAGE -> BASE64
# ============================================================

def image_to_data_uri(data, filename):

    ext = Path(filename).suffix.lower()

    mime_map = {
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".png": "image/png",
        ".gif": "image/gif",
        ".webp": "image/webp",
        ".bmp": "image/bmp",
    }

    mime = mime_map.get(
        ext,
        mimetypes.guess_type(filename)[0]
        or "application/octet-stream"
    )

    encoded = base64.b64encode(data).decode("ascii")

    return f"data:{mime};base64,{encoded}"


# ============================================================
# NORMALISASI PATH XML
# ============================================================

def resolve_target(source_file, target):

    source_dir = Path(source_file).parent

    target_path = source_dir / target

    parts = []

    for part in target_path.parts:

        if part == "..":

            if parts:
                parts.pop()

        elif part not in (".", ""):

            parts.append(part)

    return "/".join(parts)


# ============================================================
# CARI DRAWING UNTUK SETIAP SHEET
# ============================================================

def get_sheet_drawings(z):

    workbook_xml = ET.fromstring(
        z.read("xl/workbook.xml")
    )

    workbook_rels = ET.fromstring(
        z.read("xl/_rels/workbook.xml.rels")
    )

    rel_map = {
        item.attrib["Id"]: item.attrib["Target"]
        for item in workbook_rels
    }

    result = {}

    sheets = workbook_xml.find("main:sheets", NS)

    for sheet in sheets:

        sheet_name = sheet.attrib["name"]

        rid = sheet.attrib[RID_ATTR]

        target = rel_map[rid]

        if target.startswith("xl/"):
            sheet_xml = target
        else:
            sheet_xml = "xl/" + target

        rel_file = (
            "xl/worksheets/_rels/"
            + Path(sheet_xml).name
            + ".rels"
        )

        if rel_file not in z.namelist():
            continue

        sheet_rels = ET.fromstring(
            z.read(rel_file)
        )

        for relation in sheet_rels:

            relation_type = relation.attrib.get(
                "Type", ""
            )

            if relation_type.endswith("/drawing"):

                drawing_path = resolve_target(
                    sheet_xml,
                    relation.attrib["Target"]
                )

                result[sheet_xml] = drawing_path

    return result


# ============================================================
# AMBIL SEMUA GAMBAR DARI DRAWING
#
# PENTING:
# Excel menyimpan gambar dalam beberapa bentuk:
#
# 1. xdr:pic
# 2. xdr:sp + a:blipFill
# 3. twoCellAnchor
# 4. oneCellAnchor
#
# Script ini membaca SEMUANYA.
# ============================================================

def extract_images_from_drawing(z, drawing_path):

    if not drawing_path:
        return {}

    if drawing_path not in z.namelist():
        return {}

    drawing_xml = ET.fromstring(
        z.read(drawing_path)
    )

    rel_file = (
        "xl/drawings/_rels/"
        + Path(drawing_path).name
        + ".rels"
    )

    if rel_file not in z.namelist():
        return {}

    drawing_rels = ET.fromstring(
        z.read(rel_file)
    )

    relation_map = {
        item.attrib["Id"]: item.attrib["Target"]
        for item in drawing_rels
    }

    images_by_row = {}

    # Semua anchor:
    #
    # twoCellAnchor
    # oneCellAnchor
    #
    for anchor in drawing_xml:

        from_element = anchor.find(
            "xdr:from",
            NS
        )

        if from_element is None:
            continue

        row_element = from_element.find(
            "xdr:row",
            NS
        )

        col_element = from_element.find(
            "xdr:col",
            NS
        )

        if row_element is None or col_element is None:
            continue

        row_from = int(row_element.text) + 1

        col_from = int(col_element.text) + 1

        # Kolom gambar berada di sekitar E/F.
        # Ada beberapa gambar yang mulai dari E,
        # jadi jangan hanya menggunakan F.
        if col_from < 5:
            continue

        # ----------------------------------------------------
        # TENTUKAN BARIS GAMBAR
        # ----------------------------------------------------

        to_element = anchor.find(
            "xdr:to",
            NS
        )

        if to_element is not None:

            row_element_to = to_element.find(
                "xdr:row",
                NS
            )

            if row_element_to is not None:

                row_to = int(
                    row_element_to.text
                ) + 1

                # Kalau gambar hanya berada pada satu baris
                # gunakan row_from.
                #
                # Kalau gambar membentang beberapa baris,
                # gunakan posisi tengah.
                if row_to > row_from:
                    target_row = (
                        row_from + row_to
                    ) // 2
                else:
                    target_row = row_from

            else:
                target_row = row_from

        else:

            # oneCellAnchor tidak mempunyai <xdr:to>
            # sehingga gunakan row_from.
            target_row = row_from

        # ----------------------------------------------------
        # CARI SEMUA <a:blip>
        #
        # Ini yang memperbaiki masalah utama.
        # ----------------------------------------------------

        media_paths = []

        blips = anchor.findall(
            ".//a:blip",
            NS
        )

        for blip in blips:

            rid = blip.attrib.get(
                EMBED_ATTR
            )

            if not rid:
                continue

            if rid not in relation_map:
                continue

            target = relation_map[rid]

            media_path = resolve_target(
                drawing_path,
                target
            )

            if not media_path.startswith(
                "xl/media/"
            ):
                continue

            if media_path not in media_paths:
                media_paths.append(
                    media_path
                )

        if not media_paths:
            continue

        # ----------------------------------------------------
        # SIMPAN GAMBAR BERDASARKAN BARIS
        # ----------------------------------------------------

        if target_row not in images_by_row:

            images_by_row[target_row] = []

        for media_path in media_paths:

            if media_path not in images_by_row[target_row]:

                images_by_row[target_row].append(
                    media_path
                )

    return images_by_row


# ============================================================
# MULAI MEMPROSES
# ============================================================

print()
print("=" * 70)
print("MEMPROSES WEBSITE DATA INVENTARIS BNPB")
print("=" * 70)
print()


# ============================================================
# CEK FILE
# ============================================================

if not EXCEL_FILE.exists():

    print()
    print("ERROR:")
    print(
        f"File '{EXCEL_FILE.name}' tidak ditemukan."
    )
    print()
    print(
        "Pastikan file Excel berada satu folder dengan"
    )
    print("build_website.py")
    print()

    raise SystemExit


if not LOGO_FILE.exists():

    print(
        "PERINGATAN: logo_bnpb.png tidak ditemukan."
    )


# ============================================================
# BUKA EXCEL
# ============================================================

workbook = openpyxl.load_workbook(
    EXCEL_FILE,
    data_only=False
)


all_sheets = []

total_data = 0

total_with_images = 0

total_media_used = set()


# ============================================================
# BACA EXCEL + XML DRAWING
# ============================================================

with zipfile.ZipFile(
    EXCEL_FILE,
    "r"
) as z:

    sheet_drawings = get_sheet_drawings(z)

    # --------------------------------------------------------
    # DAPATKAN XML SHEET SESUAI NAMA SHEET
    # --------------------------------------------------------

    workbook_xml = ET.fromstring(
        z.read("xl/workbook.xml")
    )

    workbook_rels = ET.fromstring(
        z.read("xl/_rels/workbook.xml.rels")
    )

    workbook_rel_map = {
        item.attrib["Id"]: item.attrib["Target"]
        for item in workbook_rels
    }

    sheet_xml_by_name = {}

    for sheet in workbook_xml.find(
        "main:sheets",
        NS
    ):

        sheet_name = sheet.attrib["name"]

        rid = sheet.attrib[RID_ATTR]

        target = workbook_rel_map[rid]

        if target.startswith("xl/"):
            sheet_xml = target
        else:
            sheet_xml = "xl/" + target

        sheet_xml_by_name[sheet_name] = sheet_xml


    # ========================================================
    # PROSES SETIAP SHEET
    # ========================================================

    for worksheet in workbook.worksheets:

        print(
            f"Memproses: {worksheet.title}"
        )

        sheet_xml = sheet_xml_by_name.get(
            worksheet.title
        )

        drawing_path = sheet_drawings.get(
            sheet_xml
        )

        image_map = extract_images_from_drawing(
            z,
            drawing_path
        )

        rows = []

        # ----------------------------------------------------
        # DATA MULAI BARIS 3
        # ----------------------------------------------------

        for row_number in range(
            3,
            worksheet.max_row + 1
        ):

            nomor = worksheet.cell(
                row_number,
                1
            ).value

            nama = worksheet.cell(
                row_number,
                2
            ).value

            # Abaikan baris benar-benar kosong
            if nomor is None and nama is None:
                continue

            # ------------------------------------------------
            # AMBIL GAMBAR
            # ------------------------------------------------

            image_data = []

            media_paths = image_map.get(
                row_number,
                []
            )

            for media_path in media_paths:

                if media_path not in z.namelist():
                    continue

                try:

                    raw_image = z.read(
                        media_path
                    )

                    data_uri = image_to_data_uri(
                        raw_image,
                        media_path
                    )

                    image_data.append(
                        data_uri
                    )

                    total_media_used.add(
                        media_path
                    )

                except Exception as error:

                    print(
                        "Gagal membaca:",
                        media_path,
                        error
                    )

            # ------------------------------------------------
            # MASUKKAN DATA
            # ------------------------------------------------

            row_data = {

                "no": clean(
                    worksheet.cell(
                        row_number,
                        1
                    ).value
                ),

                "nama": clean(
                    worksheet.cell(
                        row_number,
                        2
                    ).value
                ),

                "merek": clean(
                    worksheet.cell(
                        row_number,
                        3
                    ).value
                ),

                "satuan": clean(
                    worksheet.cell(
                        row_number,
                        4
                    ).value
                ),

                "pengajuan": clean(
                    worksheet.cell(
                        row_number,
                        5
                    ).value
                ),

                "waktu": clean(
                    worksheet.cell(
                        row_number,
                        8
                    ).value
                ),

                "gambar": image_data,
            }

            rows.append(row_data)

        jumlah_bergambar = sum(
            1
            for item in rows
            if item["gambar"]
        )

        total_data += len(rows)

        total_with_images += jumlah_bergambar

        print(
            f"  Data          : {len(rows)}"
        )

        print(
            f"  Dengan gambar : {jumlah_bergambar}"
        )

        print(
            f"  Tanpa gambar  : {len(rows) - jumlah_bergambar}"
        )

        print()

        all_sheets.append(
            {
                "name": worksheet.title,
                "rows": rows
            }
        )


# ============================================================
# LOGO BNPB
# ============================================================

logo_data = ""

if LOGO_FILE.exists():

    logo_data = image_to_data_uri(
        LOGO_FILE.read_bytes(),
        LOGO_FILE.name
    )


# ============================================================
# UBAH DATA MENJADI JSON
# ============================================================

json_data = json.dumps(
    all_sheets,
    ensure_ascii=False,
    separators=(",", ":")
)

logo_json = json.dumps(
    logo_data
)


# ============================================================
# HTML WEBSITE
# ============================================================

HTML = r'''
<!DOCTYPE html>

<html lang="id">

<head>

<meta charset="UTF-8">

<meta
    name="viewport"
    content="width=device-width, initial-scale=1.0"
>

<title>
Data Inventaris Spare Part - BNPB
</title>


<style>

:root {

    --blue:
        #0B4778;

    --blue-dark:
        #07385F;

    --orange:
        #F47721;

    --background:
        #EEF3F7;

    --white:
        #FFFFFF;

    --border:
        #D9E2EA;

    --text:
        #193047;

}


* {

    box-sizing:
        border-box;

}


body {

    margin:
        0;

    font-family:
        "Segoe UI",
        Arial,
        sans-serif;

    background:
        var(--background);

    color:
        var(--text);

}


/* =========================================================
   HEADER
   ========================================================= */

header {

    background:
        linear-gradient(
            135deg,
            var(--blue-dark),
            var(--blue)
        );

    color:
        white;

    padding:
        18px 30px;

    display:
        flex;

    align-items:
        center;

    gap:
        18px;

    border-bottom:
        7px solid var(--orange);

}


.logo {

    width:
        82px;

    height:
        82px;

    background:
        white;

    padding:
        7px;

    border-radius:
        12px;

    object-fit:
        contain;

}


.title h1 {

    margin:
        0;

    font-size:
        28px;

}


.title p {

    margin:
        7px 0 0;

    font-size:
        15px;

}


/* =========================================================
   CONTAINER
   ========================================================= */

.container {

    max-width:
        1550px;

    margin:
        24px auto;

    padding:
        0 20px;

}


/* =========================================================
   STATISTIC CARDS
   ========================================================= */

.cards {

    display:
        grid;

    grid-template-columns:
        repeat(4, 1fr);

    gap:
        16px;

    margin-bottom:
        18px;

}


.card {

    background:
        white;

    border:
        1px solid var(--border);

    border-radius:
        15px;

    padding:
        19px 21px;

    box-shadow:
        0 4px 15px rgba(
            21,
            52,
            75,
            0.08
        );

}


.card-label {

    font-size:
        13px;

    font-weight:
        700;

    color:
        #64788C;

    text-transform:
        uppercase;

}


.card-value {

    font-size:
        28px;

    font-weight:
        800;

    margin-top:
        7px;

    color:
        var(--blue-dark);

}


/* =========================================================
   PANEL
   ========================================================= */

.panel {

    background:
        white;

    border:
        1px solid var(--border);

    border-radius:
        15px;

    overflow:
        hidden;

    box-shadow:
        0 4px 15px rgba(
            21,
            52,
            75,
            0.06
        );

}


/* =========================================================
   SEARCH
   ========================================================= */

.toolbar {

    padding:
        17px;

    display:
        flex;

    gap:
        12px;

    border-bottom:
        1px solid var(--border);

}


#search {

    flex:
        1;

}


input,
select {

    height:
        46px;

    border:
        1px solid #CCD8E2;

    border-radius:
        10px;

    padding:
        0 14px;

    font-size:
        14px;

}


select {

    min-width:
        150px;

}


/* =========================================================
   TABS
   ========================================================= */

.tabs {

    display:
        flex;

    gap:
        7px;

    padding:
        12px 15px;

    border-bottom:
        1px solid var(--border);

    overflow-x:
        auto;

}


.tab {

    border:
        0;

    background:
        #EDF3F7;

    color:
        #47627A;

    padding:
        11px 16px;

    border-radius:
        9px;

    font-weight:
        700;

    cursor:
        pointer;

    white-space:
        nowrap;

}


.tab.active {

    background:
        var(--blue);

    color:
        white;

}


/* =========================================================
   INFO
   ========================================================= */

.info {

    padding:
        12px 17px;

    color:
        #687B8D;

    font-size:
        13px;

}


.count {

    font-weight:
        700;

    color:
        var(--blue);

}


/* =========================================================
   TABLE
   ========================================================= */

.table-wrap {

    overflow:
        auto;

    max-height:
        68vh;

}


table {

    width:
        100%;

    border-collapse:
        separate;

    border-spacing:
        0;

    min-width:
        1050px;

}


th {

    position:
        sticky;

    top:
        0;

    background:
        var(--blue-dark);

    color:
        white;

    padding:
        13px 12px;

    text-align:
        left;

    font-size:
        13px;

    z-index:
        2;

}


td {

    padding:
        10px 12px;

    border-bottom:
        1px solid #E5EBEF;

    vertical-align:
        middle;

    font-size:
        14px;

}


tbody tr:nth-child(even) {

    background:
        #F8FAFC;

}


tbody tr:hover {

    background:
        #EEF6FB;

}


.num {

    width:
        65px;

    text-align:
        center;

}


.img-cell {

    width:
        160px;

    text-align:
        center;

}


/* =========================================================
   GAMBAR
   ========================================================= */

.part-image {

    width:
        110px;

    height:
        85px;

    object-fit:
        contain;

    background:
        white;

    border:
        1px solid #D8E1E8;

    border-radius:
        8px;

    padding:
        4px;

    cursor:
        zoom-in;

}


.image-count {

    font-size:
        11px;

    color:
        #667B8D;

    margin-top:
        3px;

}


.no-image {

    width:
        110px;

    height:
        85px;

    border:
        1px dashed #C3CED8;

    border-radius:
        8px;

    display:
        inline-flex;

    align-items:
        center;

    justify-content:
        center;

    text-align:
        center;

    color:
        #7C8B98;

    font-size:
        11px;

    padding:
        8px;

    background:
        #F8FAFC;

}


/* =========================================================
   MODAL GAMBAR
   ========================================================= */

.modal {

    display:
        none;

    position:
        fixed;

    inset:
        0;

    background:
        rgba(
            0,
            17,
            29,
            0.85
        );

    z-index:
        999;

    align-items:
        center;

    justify-content:
        center;

    padding:
        25px;

}


.modal.show {

    display:
        flex;

}


.modal-box {

    max-width:
        92vw;

    max-height:
        92vh;

    background:
        white;

    border-radius:
        14px;

    padding:
        14px;

    position:
        relative;

}


.modal-box img {

    max-width:
        88vw;

    max-height:
        84vh;

    object-fit:
        contain;

}


.close {

    position:
        absolute;

    right:
        8px;

    top:
        7px;

    border:
        0;

    background:
        white;

    border-radius:
        50%;

    width:
        32px;

    height:
        32px;

    font-size:
        20px;

    cursor:
        pointer;

    box-shadow:
        0 2px 8px rgba(
            0,
            0,
            0,
            0.2
        );

}


.empty {

    padding:
        35px;

    text-align:
        center;

    color:
        #778899;

}


footer {

    text-align:
        center;

    color:
        #718396;

    font-size:
        13px;

    padding:
        24px;

}


@media(max-width:900px) {

    .cards {

        grid-template-columns:
            repeat(2, 1fr);

    }

}


@media(max-width:560px) {

    .cards {

        grid-template-columns:
            1fr;

    }

    header {

        padding:
            15px;

    }

    .logo {

        width:
            65px;

        height:
            65px;

    }

    .title h1 {

        font-size:
            20px;

    }

}

</style>

</head>


<body>


<header>

    <img
        id="logo"
        class="logo"
        alt="Logo BNPB"
    >

    <div class="title">

        <h1>
            Data Inventaris Spare Part
        </h1>

        <p>
            Badan Nasional Penanggulangan Bencana (BNPB)
        </p>

    </div>

</header>


<div class="container">


    <!-- STATISTIK -->

    <div class="cards">

        <div class="card">

            <div class="card-label">
                Total Spare Part
            </div>

            <div
                class="card-value"
                id="total"
            >
                0
            </div>

        </div>


        <div class="card">

            <div class="card-label">
                Total Lokasi
            </div>

            <div
                class="card-value"
                id="lokasi"
            >
                0
            </div>

        </div>


        <div class="card">

            <div class="card-label">
                Data Dengan Gambar
            </div>

            <div
                class="card-value"
                id="ada"
            >
                0
            </div>

        </div>


        <div class="card">

            <div class="card-label">
                Data Tanpa Gambar
            </div>

            <div
                class="card-value"
                id="tidak"
            >
                0
            </div>

        </div>

    </div>


    <!-- PANEL -->

    <div class="panel">


        <!-- SEARCH -->

        <div class="toolbar">

            <input
                id="search"
                type="text"
                placeholder="Cari nama spare part, merk, nomor, atau tahun..."
            >

            <select id="year">

                <option value="">
                    Semua Tahun
                </option>

            </select>

        </div>


        <!-- TAB 4 SHEET -->

        <div
            class="tabs"
            id="tabs"
        ></div>


        <div class="info">

            <span id="sheetTitle"></span>

            ·

            <span
                id="resultCount"
                class="count"
            >
                0
            </span>

            data

        </div>


        <!-- TABLE -->

        <div class="table-wrap">

            <table>

                <thead>

                    <tr>

                        <th class="num">
                            NO.
                        </th>

                        <th>
                            NAMA PART
                        </th>

                        <th>
                            MERK
                        </th>

                        <th>
                            SATUAN
                        </th>

                        <th>
                            PENGAJUAN
                        </th>

                        <th class="img-cell">
                            GAMBAR
                        </th>

                        <th>
                            WAKTU PEMESANAN
                        </th>

                    </tr>

                </thead>


                <tbody id="tableBody">

                </tbody>

            </table>

        </div>

    </div>


    <footer>

        Sistem Informasi Data Inventaris Spare Part — BNPB

    </footer>

</div>


<!-- MODAL -->

<div
    class="modal"
    id="modal"
    onclick="closeModal(event)"
>

    <div class="modal-box">

        <button
            class="close"
            onclick="hideModal()"
        >
            ×
        </button>

        <img
            id="modalImage"
            alt="Gambar Spare Part"
        >

    </div>

</div>


<script>


// ============================================================
// DATA DARI PYTHON
// ============================================================

const DATA = __DATA__;

const LOGO = __LOGO__;


// ============================================================
// ELEMENT
// ============================================================

const logo =
    document.getElementById("logo");

const total =
    document.getElementById("total");

const lokasi =
    document.getElementById("lokasi");

const ada =
    document.getElementById("ada");

const tidak =
    document.getElementById("tidak");

const search =
    document.getElementById("search");

const year =
    document.getElementById("year");

const tabs =
    document.getElementById("tabs");

const tableBody =
    document.getElementById("tableBody");

const sheetTitle =
    document.getElementById("sheetTitle");

const resultCount =
    document.getElementById("resultCount");

const modal =
    document.getElementById("modal");

const modalImage =
    document.getElementById("modalImage");


logo.src = LOGO;


// ============================================================
// DATA GLOBAL
// ============================================================

const ALL_DATA =
    DATA.flatMap(
        sheet => sheet.rows
    );


total.textContent =
    ALL_DATA.length;


lokasi.textContent =
    DATA.length;


const totalWithImage =
    ALL_DATA.filter(
        item =>
            item.gambar &&
            item.gambar.length > 0
    ).length;


ada.textContent =
    totalWithImage;


tidak.textContent =
    ALL_DATA.length -
    totalWithImage;


// ============================================================
// TAHUN
// ============================================================

const years = [
    ...new Set(
        ALL_DATA.flatMap(
            item =>
                (
                    item.waktu.match(
                        /20\d{2}/g
                    ) || []
                )
        )
    )
].sort();


years.forEach(
    item => {

        const option =
            document.createElement(
                "option"
            );

        option.value =
            item;

        option.textContent =
            item;

        year.appendChild(
            option
        );

    }
);


// ============================================================
// TAB
// ============================================================

let activeSheet = 0;


DATA.forEach(
    (sheet, index) => {

        const button =
            document.createElement(
                "button"
            );

        button.className =
            "tab" +
            (
                index === 0
                    ? " active"
                    : ""
            );

        button.textContent =
            sheet.name.replace(
                "Sparepart ",
                ""
            );


        button.onclick =
            function() {

                activeSheet =
                    index;

                document
                    .querySelectorAll(
                        ".tab"
                    )
                    .forEach(
                        tab =>
                            tab.classList.remove(
                                "active"
                            )
                    );

                button.classList.add(
                    "active"
                );

                render();

            };


        tabs.appendChild(
            button
        );

    }
);


// ============================================================
// ESCAPE HTML
// ============================================================

function escapeHTML(value) {

    return String(
        value ?? ""
    ).replace(
        /[&<>"']/g,
        char => ({
            "&": "&amp;",
            "<": "&lt;",
            ">": "&gt;",
            '"': "&quot;",
            "'": "&#39;"
        })[char]
    );

}


// ============================================================
// RENDER TABEL
// ============================================================

function render() {

    const query =
        search.value
            .toLowerCase()
            .trim();

    const selectedYear =
        year.value;

    const sheet =
        DATA[activeSheet];


    const filteredRows =
        sheet.rows.filter(
            item => {

                const text = [

                    item.no,
                    item.nama,
                    item.merek,
                    item.satuan,
                    item.pengajuan,
                    item.waktu

                ]
                    .join(" ")
                    .toLowerCase();


                const matchSearch =
                    !query ||
                    text.includes(
                        query
                    );


                const matchYear =
                    !selectedYear ||
                    item.waktu.includes(
                        selectedYear
                    );


                return (
                    matchSearch &&
                    matchYear
                );

            }
        );


    sheetTitle.textContent =
        sheet.name;


    resultCount.textContent =
        filteredRows.length;


    if (
        filteredRows.length === 0
    ) {

        tableBody.innerHTML = `

            <tr>

                <td
                    colspan="7"
                    class="empty"
                >

                    Tidak ada data yang sesuai.

                </td>

            </tr>

        `;

        return;

    }


    tableBody.innerHTML =
        filteredRows
            .map(
                item => {

                    let imagesHTML = "";


                    if (
                        item.gambar &&
                        item.gambar.length
                    ) {

                        imagesHTML =
                            item.gambar
                                .map(
                                    src => `

                                    <img
                                        class="part-image"
                                        src="${src}"
                                        onclick="showModal(this.src)"
                                        alt="Gambar ${escapeHTML(item.nama)}"
                                        title="Klik untuk memperbesar"
                                    >

                                    `
                                )
                                .join("");


                        if (
                            item.gambar.length > 1
                        ) {

                            imagesHTML += `

                                <div class="image-count">

                                    ${item.gambar.length}
                                    gambar

                                </div>

                            `;

                        }

                    }
                    else {

                        imagesHTML = `

                            <span class="no-image">

                                Gambar belum tersedia
                                pada file Excel

                            </span>

                        `;

                    }


                    return `

                        <tr>

                            <td class="num">

                                ${escapeHTML(
                                    item.no
                                )}

                            </td>


                            <td>

                                ${escapeHTML(
                                    item.nama
                                ).replace(
                                    /\n/g,
                                    "<br>"
                                )}

                            </td>


                            <td>

                                ${escapeHTML(
                                    item.merek
                                )}

                            </td>


                            <td>

                                ${escapeHTML(
                                    item.satuan
                                )}

                            </td>


                            <td>

                                ${escapeHTML(
                                    item.pengajuan
                                )}

                            </td>


                            <td class="img-cell">

                                ${imagesHTML}

                            </td>


                            <td>

                                ${escapeHTML(
                                    item.waktu
                                )}

                            </td>

                        </tr>

                    `;

                }
            )
            .join("");

}


// ============================================================
// MODAL GAMBAR
// ============================================================

function showModal(src) {

    modalImage.src =
        src;

    modal.classList.add(
        "show"
    );

}


function hideModal() {

    modal.classList.remove(
        "show"
    );

    modalImage.src =
        "";

}


function closeModal(event) {

    if (
        event.target === modal
    ) {

        hideModal();

    }

}


// ============================================================
// EVENT
// ============================================================

search.addEventListener(
    "input",
    render
);


year.addEventListener(
    "change",
    render
);


// ============================================================
// TAMPILKAN AWAL
// ============================================================

render();


</script>

</body>

</html>
'''


# ============================================================
# MASUKKAN DATA KE HTML
# ============================================================

HTML = HTML.replace(
    "__DATA__",
    json_data
)

HTML = HTML.replace(
    "__LOGO__",
    logo_json
)


# ============================================================
# SIMPAN
# ============================================================

OUTPUT_FILE.write_text(
    HTML,
    encoding="utf-8"
)


# ============================================================
# HASIL
# ============================================================

print()
print("=" * 70)
print("WEBSITE BERHASIL DIBUAT")
print("=" * 70)

print(
    "File              :",
    OUTPUT_FILE
)

print(
    "Jumlah sheet      :",
    len(all_sheets)
)

print(
    "Total data        :",
    total_data
)

print(
    "Data dengan gambar:",
    total_with_images
)

print(
    "Data tanpa gambar :",
    total_data - total_with_images
)

print(
    "Media gambar dipakai:",
    len(total_media_used)
)

print("=" * 70)
print()