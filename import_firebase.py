import json
import mimetypes
import re
import zipfile
from pathlib import Path
import xml.etree.ElementTree as ET

import openpyxl
import firebase_admin
from firebase_admin import credentials, db


# ============================================================
# KONFIGURASI
# ============================================================

BASE_DIR = Path(__file__).resolve().parent

EXCEL_FILE = BASE_DIR / "SPAREPART BNPB.xlsx"

# Folder gambar akan dibuat di sebelah index.html
IMAGE_DIR = BASE_DIR / "images"

# Firebase Realtime Database
DATABASE_URL = (
    "https://inventaris-sparepart-bnpb-default-rtdb.firebaseio.com/"
)

# Nama folder di Firebase
FIREBASE_ROOT = "inventaris"

# Nama sheet Excel -> nama folder di Firebase / folder gambar
LOCATION_MAP = {
    "Sentul": "sentul",
    "Jatiasih": "jatiasih",
    "Pramuka": "pramuka",
    "UPT Padang": "upt_padang",
}


# ============================================================
# XML NAMESPACE
# ============================================================

NS = {
    "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "rel": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}

RID_ATTR = (
    "{http://schemas.openxmlformats.org/officeDocument/"
    "2006/relationships}id"
)

EMBED_ATTR = (
    "{http://schemas.openxmlformats.org/officeDocument/"
    "2006/relationships}embed"
)


# ============================================================
# MEMBERSIHKAN DATA
# ============================================================

def clean(value):
    if value is None:
        return ""

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    return str(value).strip()


# ============================================================
# NORMALISASI PATH XML
# ============================================================

def resolve_target(source_file, target):
    source_dir = Path(source_file).parent
    target_path = source_dir / target

    parts = []

    for part in target_path.parts:
        if part == "..":
            if parts:
                parts.pop()
        elif part not in (".", ""):
            parts.append(part)

    return "/".join(parts)


# ============================================================
# CARI DRAWING UNTUK SETIAP SHEET
# ============================================================

def get_sheet_drawings(z):
    workbook_xml = ET.fromstring(
        z.read("xl/workbook.xml")
    )

    workbook_rels = ET.fromstring(
        z.read("xl/_rels/workbook.xml.rels")
    )

    rel_map = {
        item.attrib["Id"]: item.attrib["Target"]
        for item in workbook_rels
    }

    result = {}

    sheets = workbook_xml.find("main:sheets", NS)

    for sheet in sheets:
        sheet_name = sheet.attrib["name"]
        rid = sheet.attrib[RID_ATTR]
        target = rel_map[rid]

        if target.startswith("xl/"):
            sheet_xml = target
        else:
            sheet_xml = "xl/" + target

        rel_file = (
            "xl/worksheets/_rels/"
            + Path(sheet_xml).name
            + ".rels"
        )

        if rel_file not in z.namelist():
            continue

        sheet_rels = ET.fromstring(
            z.read(rel_file)
        )

        for relation in sheet_rels:
            relation_type = relation.attrib.get("Type", "")

            if relation_type.endswith("/drawing"):
                drawing_path = resolve_target(
                    sheet_xml,
                    relation.attrib["Target"]
                )

                result[sheet_xml] = drawing_path

    return result


# ============================================================
# AMBIL GAMBAR DARI DRAWING EXCEL
# ============================================================

def extract_images_from_drawing(z, drawing_path):
    if not drawing_path:
        return {}

    if drawing_path not in z.namelist():
        return {}

    drawing_xml = ET.fromstring(
        z.read(drawing_path)
    )

    rel_file = (
        "xl/drawings/_rels/"
        + Path(drawing_path).name
        + ".rels"
    )

    if rel_file not in z.namelist():
        return {}

    drawing_rels = ET.fromstring(
        z.read(rel_file)
    )

    relation_map = {
        item.attrib["Id"]: item.attrib["Target"]
        for item in drawing_rels
    }

    images_by_row = {}

    for anchor in drawing_xml:
        from_element = anchor.find("xdr:from", NS)

        if from_element is None:
            continue

        row_element = from_element.find("xdr:row", NS)
        col_element = from_element.find("xdr:col", NS)

        if row_element is None or col_element is None:
            continue

        row_from = int(row_element.text) + 1
        col_from = int(col_element.text) + 1

        # Gambar spare part pada file Excel berada
        # di sekitar kolom E/F.
        if col_from < 5:
            continue

        to_element = anchor.find("xdr:to", NS)

        if to_element is not None:
            row_element_to = to_element.find("xdr:row", NS)

            if row_element_to is not None:
                row_to = int(row_element_to.text) + 1

                if row_to > row_from:
                    target_row = (row_from + row_to) // 2
                else:
                    target_row = row_from
            else:
                target_row = row_from
        else:
            target_row = row_from

        media_paths = []

        for blip in anchor.findall(".//a:blip", NS):
            rid = blip.attrib.get(EMBED_ATTR)

            if not rid:
                continue

            if rid not in relation_map:
                continue

            target = relation_map[rid]

            media_path = resolve_target(
                drawing_path,
                target
            )

            if not media_path.startswith("xl/media/"):
                continue

            if media_path not in media_paths:
                media_paths.append(media_path)

        if not media_paths:
            continue

        if target_row not in images_by_row:
            images_by_row[target_row] = []

        for media_path in media_paths:
            if media_path not in images_by_row[target_row]:
                images_by_row[target_row].append(media_path)

    return images_by_row


# ============================================================
# NAMA FILE GAMBAR YANG AMAN
# ============================================================

def safe_filename(text):
    text = clean(text)

    text = re.sub(
        r'[<>:"/\\|?*\x00-\x1F]',
        "_",
        text
    )

    text = re.sub(
        r"\s+",
        "_",
        text
    )

    text = text.strip("._ ")

    if not text:
        text = "sparepart"

    return text[:100]


# ============================================================
# EKSTENSI GAMBAR
# ============================================================

def image_extension(media_path):
    ext = Path(media_path).suffix.lower()

    if ext in {
        ".jpg",
        ".jpeg",
        ".png",
        ".gif",
        ".webp",
        ".bmp",
        ".svg",
    }:
        return ext

    mime = mimetypes.guess_type(media_path)[0] or ""

    mime_map = {
        "image/jpeg": ".jpg",
        "image/png": ".png",
        "image/gif": ".gif",
        "image/webp": ".webp",
        "image/bmp": ".bmp",
        "image/svg+xml": ".svg",
    }

    return mime_map.get(mime, ".png")


# ============================================================
# INISIALISASI FIREBASE
# ============================================================

def find_service_account():
    files = sorted(
        BASE_DIR.glob("*firebase-adminsdk*.json")
    )

    if not files:
        return None

    return files[0]


def init_firebase():
    service_account = find_service_account()

    if service_account is None:
        print()
        print("ERROR: File Firebase Admin SDK tidak ditemukan.")
        print()
        print(
            "Pastikan file JSON service account berada satu folder"
        )
        print("dengan import_firebase.py.")
        print()
        raise SystemExit(1)

    print(
        "Service account :",
        service_account.name
    )

    if not firebase_admin._apps:
        cred = credentials.Certificate(
            str(service_account)
        )

        firebase_admin.initialize_app(
            cred,
            {
                "databaseURL": DATABASE_URL
            }
        )

    return db.reference(FIREBASE_ROOT)


# ============================================================
# MULAI
# ============================================================

print()
print("=" * 70)
print("IMPORT DATA BNPB -> FIREBASE REALTIME DATABASE")
print("GAMBAR -> FOLDER WEBSITE (TANPA FIREBASE STORAGE)")
print("=" * 70)
print()


# ============================================================
# CEK EXCEL
# ============================================================

if not EXCEL_FILE.exists():
    print("ERROR:")
    print(
        f"File '{EXCEL_FILE.name}' tidak ditemukan."
    )
    print()
    print(
        "Pastikan file Excel berada satu folder dengan"
    )
    print("import_firebase.py.")
    print()
    raise SystemExit(1)


# ============================================================
# SIAPKAN FOLDER GAMBAR
# ============================================================

IMAGE_DIR.mkdir(
    parents=True,
    exist_ok=True
)


# ============================================================
# FIREBASE
# ============================================================

firebase_root = init_firebase()


# ============================================================
# BUKA EXCEL
# ============================================================

workbook = openpyxl.load_workbook(
    EXCEL_FILE,
    data_only=False
)


total_data = 0
total_with_images = 0
total_images_saved = 0


# ============================================================
# BACA XML EXCEL
# ============================================================

with zipfile.ZipFile(
    EXCEL_FILE,
    "r"
) as z:

    sheet_drawings = get_sheet_drawings(z)

    workbook_xml = ET.fromstring(
        z.read("xl/workbook.xml")
    )

    workbook_rels = ET.fromstring(
        z.read("xl/_rels/workbook.xml.rels")
    )

    workbook_rel_map = {
        item.attrib["Id"]: item.attrib["Target"]
        for item in workbook_rels
    }

    sheet_xml_by_name = {}

    for sheet in workbook_xml.find(
        "main:sheets",
        NS
    ):
        sheet_name = sheet.attrib["name"]

        rid = sheet.attrib[RID_ATTR]

        target = workbook_rel_map[rid]

        if target.startswith("xl/"):
            sheet_xml = target
        else:
            sheet_xml = "xl/" + target

        sheet_xml_by_name[sheet_name] = sheet_xml


    # ========================================================
    # PROSES SETIAP SHEET
    # ========================================================

    for worksheet in workbook.worksheets:

        sheet_name = worksheet.title

        location = LOCATION_MAP.get(
            sheet_name,
            safe_filename(sheet_name).lower()
        )

        print()
        print("-" * 70)
        print(f"SHEET    : {sheet_name}")
        print(f"FIREBASE : {FIREBASE_ROOT}/{location}")
        print("-" * 70)

        sheet_xml = sheet_xml_by_name.get(
            sheet_name
        )

        drawing_path = sheet_drawings.get(
            sheet_xml
        )

        image_map = extract_images_from_drawing(
            z,
            drawing_path
        )

        location_image_dir = (
            IMAGE_DIR / location
        )

        location_image_dir.mkdir(
            parents=True,
            exist_ok=True
        )

        location_ref = firebase_root.child(
            location
        )

        rows_imported = 0
        rows_with_images = 0
        images_saved_this_sheet = 0

        # ----------------------------------------------------
        # DATA MULAI BARIS 3
        # ----------------------------------------------------

        for row_number in range(
            3,
            worksheet.max_row + 1
        ):

            nomor = worksheet.cell(
                row_number,
                1
            ).value

            nama = worksheet.cell(
                row_number,
                2
            ).value

            # Abaikan baris kosong
            if nomor is None and nama is None:
                continue

            # ------------------------------------------------
            # DATA EXCEL
            # ------------------------------------------------

            no_value = clean(
                worksheet.cell(
                    row_number,
                    1
                ).value
            )

            nama_value = clean(
                worksheet.cell(
                    row_number,
                    2
                ).value
            )

            merek_value = clean(
                worksheet.cell(
                    row_number,
                    3
                ).value
            )

            satuan_value = clean(
                worksheet.cell(
                    row_number,
                    4
                ).value
            )

            pengajuan_value = clean(
                worksheet.cell(
                    row_number,
                    5
                ).value
            )

            waktu_value = clean(
                worksheet.cell(
                    row_number,
                    8
                ).value
            )

            # ------------------------------------------------
            # GAMBAR
            # ------------------------------------------------

            image_urls = []

            media_paths = image_map.get(
                row_number,
                []
            )

            for image_index, media_path in enumerate(
                media_paths,
                start=1
            ):

                if media_path not in z.namelist():
                    continue

                try:
                    raw_image = z.read(
                        media_path
                    )

                    ext = image_extension(
                        media_path
                    )

                    filename = (
                        f"row_{row_number}_"
                        f"{image_index}"
                        f"{ext}"
                    )

                    image_path = (
                        location_image_dir
                        / filename
                    )

                    image_path.write_bytes(
                        raw_image
                    )

                    # Path relatif yang bisa digunakan
                    # langsung oleh website.
                    relative_path = (
                        f"images/{location}/{filename}"
                    )

                    image_urls.append(
                        relative_path
                    )

                    images_saved_this_sheet += 1
                    total_images_saved += 1

                except Exception as error:
                    print(
                        f"  ! Gagal menyimpan gambar "
                        f"baris {row_number}: {error}"
                    )

            if image_urls:
                rows_with_images += 1
                total_with_images += 1

            # ------------------------------------------------
            # DATA UNTUK FIREBASE
            # ------------------------------------------------

            row_data = {
                "no": no_value,
                "nama": nama_value,
                "merek": merek_value,
                "satuan": satuan_value,
                "pengajuan": pengajuan_value,
                "waktu": waktu_value,
                "gambar": image_urls,
                "lokasi": sheet_name,
                "sumber": "Excel",
                "excel_row": row_number,
            }

            # Kunci dibuat dari nomor baris Excel supaya
            # stabil dan tidak membuat duplikasi saat script
            # dijalankan kembali.
            firebase_key = (
                f"excel_row_{row_number}"
            )

            location_ref.child(
                firebase_key
            ).set(row_data)

            rows_imported += 1
            total_data += 1

            print(
                f"  ✓ Data {no_value or row_number} | "
                f"{nama_value[:55]}"
            )

        print()
        print(
            f"  Data masuk Firebase : {rows_imported}"
        )
        print(
            f"  Dengan gambar       : {rows_with_images}"
        )
        print(
            f"  File gambar dibuat  : {images_saved_this_sheet}"
        )


# ============================================================
# HASIL AKHIR
# ============================================================

print()
print("=" * 70)
print("IMPORT SELESAI")
print("=" * 70)
print()
print(
    "Total data Firebase       :",
    total_data
)
print(
    "Data dengan gambar        :",
    total_with_images
)
print(
    "File gambar dibuat        :",
    total_images_saved
)
print(
    "Folder gambar             :",
    IMAGE_DIR
)
print()
print(
    "Firebase Storage TIDAK digunakan."
)
print(
    "Firebase Realtime Database tetap digunakan."
)
print()
print(
    "Langkah berikutnya:"
)
print(
    "Pastikan folder 'images' ikut di-upload ke GitHub"
)
print(
    "bersama index.html."
)
print()
